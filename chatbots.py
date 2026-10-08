from openai.lib.streaming import AsyncAssistantEventHandler
from typing_extensions import override
from channels.layers import get_channel_layer
from apps.api.models import ChatBotModel,ExcelProjects,ExcelChatBotModel,GeneralChatbotModel
from apps.users.models import Files,ProjectQuery,User
from openai import AssistantEventHandler 
import queue
import os
import asyncio
from apps.api.data_classes import DocxFields
from apps.api.queries import docx_queries_instance
from threading import Thread
import functools
import re
from django.core.cache import cache
from asgiref.sync import sync_to_async
import json
from channels.db import database_sync_to_async
from io import BytesIO
import pandas as pd
import base64
from typing import TypedDict
import dropbox
from apps.helpersapps.integrations.dropbox_client import dbx, FOLDER_PATH
from apps.helpersapps.llm.clients import AZURE_CLIENT, AZURE_CLIENT_1, get_azure_async_client
from apps.helpersapps.llm.apicalling import asyncopenaiapi
from apps.helpersapps.core.aimodels import OPENAI_MODEL_5, OPENAI_MODEL_ENV1
from openai.types.beta.threads.image_file_content_block import ImageFileContentBlock
from openai.types.beta.threads.text_content_block import TextContentBlock

NON_COUNT_KEYS = {"is_research_question", "user_themes", "child"}

class ExcelChatBotPayload(TypedDict):
    image_data: str | None
    text_response: str | None
    thread_id: str | None
    assistant_id: str | None
    error: bool 
    error_message: str | None
    
def chat_uploaded_files_filter_id(file_ids):
    queryset = Files.objects.filter(id__in=file_ids).only("file")
    uploaded_files = [file.file.name for file in queryset]
    return uploaded_files

@database_sync_to_async
def get_database_instance(model_name,id):
    return model_name.objects.filter(Project = id).first()


def first_user_message_custom(project_id, name):
    try:
        user_message = ""
        obj = ProjectQuery.objects.filter(id=project_id).first()
       
        themes_datas = json.loads(obj.theme_transcript) if obj.theme_transcript else []
        def prepmesage(results):
            nonlocal user_message
            for r in results:
                doc_name = list(r.keys())[0]
                content = r[doc_name]
                user_message += f"write a synthesis on the topic '{name}' \n\n{doc_name} {content}\n\n"

        for theme_data in themes_datas:
            if theme_data.get("theme").rstrip() == name.rstrip():
                results = theme_data["results"]
                prepmesage(results)
            else:
                child = theme_data.get("child", [])
                for c in child:
                    if c.get("theme").rstrip() == name.rstrip():
                        results = c.get("results", [])
                        prepmesage(results)
                        
        if not user_message:
            user_message = name
        return user_message
    except Exception as e:
        print("An error occurred:", str(e))
        return user_message
    
def data_prep(project_id):
    MAX_LENGTH = 256000

    obj = (
        docx_queries_instance._get_project_by_id(project_id)
        .values(DocxFields.ANSWER, DocxFields.ANSWER_THEME)
        .first()
    )
    if not obj:
        return []

    chunks = []
    total = 0 

    if obj["answer"]:
        answer = json.loads(obj["answer"]).get("test")
        if answer:
            chunks.append(answer)
            total += len(answer) + 1

    answer_theme = json.loads(obj["answer_theme"]) if obj["answer_theme"] else None

    if answer_theme:
        for value in answer_theme:
            if total >= MAX_LENGTH:
                break
            theme = next((key for key in value if key not in NON_COUNT_KEYS), None)
            if theme is None:
                continue
            part = f"**{theme}**: {value[theme]}\n\n"
            chunks.append(part)
            total += len(part) + 1
    else:
        rows = (
            docx_queries_instance.get_docx_project_result_by_field(
                project=project_id, field_name=DocxFields.ANSWER_THEME
            )
            .filter(version=1)
            .values_list("data", flat=True)
            .iterator()
        )
        for row in rows:
            if total >= MAX_LENGTH:
                break
            summary = json.loads(row).get("Summary")
            if not summary:
                continue
            chunks.append(summary)
            total += len(summary) + 1

    return "\n".join(chunks)[:MAX_LENGTH]



class EventHandler(AssistantEventHandler):    
    
    def __init__(self:object, queue:queue.Queue):
        super().__init__() 
        self._queue =  queue
        self._result = None
       
    @override
    def on_text_created(self, text) -> None:
        print(f"\nassistant > ", end="", flush=True)
        
    @override
    def on_text_delta(self, delta, snapshot):
        self._queue.put(delta.value)
        
    @override
    def on_image_file_done(self, image_file):
        """Callback that is fired when an image file block is finished"""
        self._queue.put(("image_file_id", image_file.file_id))
      
    @override
    def on_message_done(self, message) -> None:
        for block in message.content:
            if isinstance(block, ImageFileContentBlock):
                print("This block is an image block.")
            elif isinstance(block, TextContentBlock):
                self._result = block.text.value
            else:
                print("This block type is unrecognized.")


    @override
    def on_event(self, event) -> None:
        if event.event == "thread.run.requires_action":
            run_id = event.data.id
            self._handle_requires_action(event.data, run_id)

    def _handle_requires_action(self, run_data, run_id: str):
        tool_outputs = []

        for tool_call in run_data.required_action.submit_tool_outputs.tool_calls:
            if tool_call.type != "function":
                continue

            if tool_call.function.name == "generate_image":
                args = json.loads(tool_call.function.arguments or "{}")
                prompt = args.get("prompt", "")

                img = AZURE_CLIENT.images.generate(
                    model="gpt-image-1",
                    prompt=prompt,
                )
                b64png = img.data[0].b64_json

            
                png_bytes = base64.b64decode(b64png)
                uploaded = AZURE_CLIENT.files.create(
                    file=("generated.png", png_bytes),
                    purpose="assistants",
                )

                tool_outputs.append({
                    "tool_call_id": tool_call.id,
                    "output": json.dumps({"image_file_id": uploaded.id}),
                })

        if tool_outputs:
            with AZURE_CLIENT.beta.threads.runs.submit_tool_outputs_stream(
                thread_id=self.current_run.thread_id,
                run_id=self.current_run.id,
                tool_outputs=tool_outputs,
                event_handler=EventHandler(self._queue),
            ) as stream:
                stream.until_done()
       


class AsyncEventHandler(AsyncAssistantEventHandler):
    result_output = ""

    def __init__(self, user, thread_id, assistant_id,message_id,user_group, azure_client=None):
        super().__init__()
        self.user = user
        self.thread_id = thread_id
        self.assistant_id = assistant_id
        self.message_id = message_id
        self.bracketed_text_list = []
        self.incomplete_bracketed_text = ""
        self.unique_bracketed_text_dict = {}
        self.user_group = user_group
        self.azure_client = azure_client

    async def send_to(self, message):
        channel_layer = get_channel_layer()
        email = self.user.email if self.user else None
        await channel_layer.group_send(
            self.user_group,
            {
                "type": "chat.message",
                "message": message,
                "user": email,
                "thread_id": self.thread_id,
                "assistant_id": self.assistant_id,
                "message_id":self.message_id
            },
        )

    @override
    async def on_text_created(self, text) -> None:
        print(f"\nassistant > ", end="", flush=True)
    
    @override
    async def on_text_delta(self, delta, snapshot):
        message = delta.value
        CITATION_RE = re.compile(r"【\d+(?::\d+)?†source】")
        message = CITATION_RE.sub("", message)
        await self.send_to(message)


    @override
    async def on_message_done(self, message) -> None:
        message_content = message.content[0].text
        annotations = message_content.annotations
        citations = []
        for index, annotation in enumerate(annotations):
            # Match the annotation.text to the stored bracketed text
            if annotation.text in self.unique_bracketed_text_dict:
                matched_index = self.unique_bracketed_text_dict[annotation.text]
                if file_citation := getattr(annotation, "file_citation", None):
                    cited_file = self.azure_client.files.retrieve(file_citation.file_id)
                    citation_text = f"[{matched_index}] {cited_file.filename}"
                    if citation_text not in citations:
                        citations.append(citation_text)

        if citations:
            my_citation = "Footnotes\n" + "\n".join(citations)
            await self.send_to(my_citation)

        message = "stream end"
        await self.send_to(message)
        AsyncEventHandler.result_output += message_content.value


# General Chatbot
class GeneralChatbotEventHandler(AssistantEventHandler):   
    def __init__(self,queue):
        super().__init__()
        self.queue = queue
        self.result = ""
   
    @override
    def on_text_delta(self, delta, snapshot):
        value = delta.value
        if value.endswith("†source】"):
            return

        self.queue.put(value)
        
    @override
    def on_message_done(self, message):
        self.result = message.content[0].text.value
        self.queue.put(None)

    @override
    def on_error(self, error):
        print("STREAM ERROR:", repr(error))
      

# Advance Analysis Docx/Pdf Chatbot..
class ChatBotHandle:
    def __init__(self:object,user:User,project_id:int,theme_type:str,user_group:str,files_attachments:list):
        self.file_ids = []
        self.channel_layer=get_channel_layer()
        self.thread_id = None
        self.assistant_id = None
        self.vector_store_id = None
        self.user = user
        self.theme_type = theme_type
        self.project_id = project_id
        self.user_group = user_group
        self.files_attachments =  files_attachments
        self._azure_client = None

    @property
    def azure_client(self):
        """Lazily create async Azure client bound to the current event loop."""
        if self._azure_client is None:
            self._azure_client = get_azure_async_client()
        return self._azure_client

    async def get_chat_name(self,assistant_response):
        if self.user:
            try:
                prompt = f""" Based on the text below, provide a descriptive chat title  Only output the title name, nothing else. \n
                    {assistant_response}
                """

                response = await asyncopenaiapi(
                    client = self.azure_client,
                    model=OPENAI_MODEL_ENV1,
                    user_message = prompt,
                    system_message = "You are a helpful assistant."

                )
               
                project_name = response.choices[0].message.content
                name = project_name.replace('"', "")
            except Exception as e:
                print("Failed while generating the title of the chat:",str(e))
                
                name = "User chat"
            await sync_to_async(ChatBotModel.objects.create)(
                    user=self.user,
                    thread_id=self.thread_id,
                    assistant_id=self.assistant_id,
                    chat_title=name,
                    theme_type=self.theme_type,
                    project_id=self.project_id,
                )


    def upload_file_to_assistant(self, response, filename):
        file_obj = BytesIO(response.content)
        file_obj.name = filename
        if file_obj:
            self.file_ids.append(file_obj)


    def download_and_extract_files(self):
        file_ids = [file['file_id'] for file in self.files_attachments]
        uploaded_files = chat_uploaded_files_filter_id(file_ids=file_ids)
        
        for i in uploaded_files:
            filename = str(i).split("/")[-1]
            b, response = dbx.files_download(path=FOLDER_PATH + str(i))
            self.upload_file_to_assistant(response, filename)
    
    
    async def prepare_prompt(self, theme_chat, project_id, name):
        if theme_chat:
            first_user_message = await asyncio.to_thread(first_user_message_custom,project_id, name)
        else:
            first_user_message = await asyncio.to_thread(data_prep, project_id)
        return first_user_message
    
    async def create_vector_store(self,name):
        try:
            if self.files_attachments:
                vector_store = await self.azure_client.vector_stores.create(
                    name=f"{name}", expires_after={"anchor": "last_active_at", "days": 100}
                )
                self.vector_store_id = vector_store.id
            
                await sync_to_async(self.download_and_extract_files)()
           
                file_batch = await self.azure_client.vector_stores.file_batches.upload_and_poll(
                    vector_store_id=vector_store.id, files=self.file_ids
                )
                print(file_batch.status)
                print(file_batch.file_counts) 
            
           
        except Exception as e:
            message = "As your file(s) have been deleted, this chatbot only has access to the analysis but will not have access to your files."
            await self.send_error_message(user_group=self.user_group,message = message)
        
        
    async def create_assistant(self):
        parameters = {
            "model":OPENAI_MODEL_ENV1,
        }
        if self.vector_store_id:
            parameters['tools'] = [{"type": "file_search"}]
            parameters['tool_resources']={"file_search": {"vector_store_ids": [self.vector_store_id]}}

        assistant = await self.azure_client.beta.assistants.create(**parameters)
        self.assistant_id = assistant.id
    
    async def create_thread(self,messages):
        thread = await self.azure_client.beta.threads.create(messages=messages)
        self.thread_id = thread.id
        
        
    async def send_error_message(self,user_group, message):
        await self.channel_layer.group_send(
            user_group,{
                "type":"trigger.error",
                "message":message
            }   
        )
        
    async def run_assistant_thread(self,thread_id,assistant_id,message_id):
        try:
            event_handler = AsyncEventHandler(self.user, thread_id,assistant_id,message_id,self.user_group, azure_client=self.azure_client)

            print("run assistant....")
            async with self.azure_client.beta.threads.runs.stream(
                thread_id=thread_id,
                assistant_id=assistant_id,
                event_handler=event_handler,
            ) as stream:
                await stream.until_done()

            print("end assistant")
            return event_handler.result_output
        except Exception as e:
            print("Error",e)
            await self.send_error_message(user_group=self.user_group,message =str(e))
        
    def save_instance(self,task):
        result = task.result()
        asyncio.create_task(self.get_chat_name(result)) 

    def perform_database_entry(self,new_thread_id,original_conversation,message_id,new_message_id):
        ChatBotModel.objects.create(
            user=original_conversation.user,
            project=original_conversation.project,
            thread_id=new_thread_id,
            assistant_id=original_conversation.assistant_id,
            parent_conversation=original_conversation,
            branch_from_message_id=message_id,
            branch_start_message_id = new_message_id,
            chat_title=original_conversation.chat_title,
            theme_type=original_conversation.theme_type
        )

        
    async def save_edited_instance(self,f,new_thread_id,original_conversation,message_id,new_message_id):
        await asyncio.to_thread(self.perform_database_entry,new_thread_id,original_conversation,message_id,new_message_id)
      
    async def newchat(self,name, theme_chat, first_ai_reply, user_message):
        await self.create_vector_store(name=name)
  
        first_user_message=await self.prepare_prompt(theme_chat=theme_chat, project_id=self.project_id, name=name)
        messages = [{"role": "user", "content": first_user_message}]

       
        if first_ai_reply:
            messages.append({"role": "assistant", "content": first_ai_reply})
        
        await self.create_assistant()
        await self.create_thread(messages=messages)
        
        message = await self.azure_client.beta.threads.messages.create(
            thread_id=self.thread_id, role="user", content=user_message
        )

        future =asyncio.create_task(self.run_assistant_thread(self.thread_id,self.assistant_id,message.id))
        future.add_done_callback(functools.partial(self.save_instance))
        
        print("start chat first message processed successfully")

    async def start_new_thread_with_history(self,messages):
        thread = await self.azure_client.beta.threads.create(messages=messages)
        return thread.id
            
    async def chat_continue(self,root_thread_id, assistant_id, user_message,chat_edit,thread_id,message_id):
        if chat_edit:
            original_conversation = await sync_to_async(ChatBotModel.objects.filter(thread_id=root_thread_id,assistant_id=assistant_id).first)()
            messages = []
            thread_messages = await self.azure_client.beta.threads.messages.list(thread_id)
            thread_messages = thread_messages.data
            thread_messages.reverse()  # Reverse to chronological order
            # Find the edit point
            edit_index = next(i for i, msg in enumerate(thread_messages) if msg.id == message_id)

            for message in thread_messages[:edit_index]:
                messages.append(
                    {
                        "content": message.content[0].text.value,
                        "role": message.role,
                    }
                )
            new_thread_id = await self.start_new_thread_with_history(messages)
           
            message = await self.azure_client.beta.threads.messages.create(
                thread_id=new_thread_id, role="user", content=user_message
            )
            future = asyncio.create_task(self.run_assistant_thread(new_thread_id,assistant_id,message.id))
            future.add_done_callback(lambda f : asyncio.create_task(self.save_edited_instance(f , new_thread_id,original_conversation,message_id,message.id)))
        else:
            message = await self.azure_client.beta.threads.messages.create(
                thread_id=thread_id, role="user", content=user_message
            )
            asyncio.create_task(self.run_assistant_thread(thread_id,assistant_id,message.id))
            print("message processed successfully")
       




# Advance Analysis Xlsx/Csv Chatbot ....
class Excelchatbot:
    def __init__(self:object, new_chat:bool,project_id:int,usermessage:str,user:str,first_ai_reply:str):
        self._new_chat:bool = new_chat
        self._project_id:int = project_id
        self._user_message:str = usermessage
        self._file:list | BytesIO | None = None
        self._user:str = user
        self._file_id: list | str | None = None
        self._assistant_id:str | None = None
        self._columns:list | None = None
        self._project = self._get_instance()
        self._user_column:str | None = self._project.column_name
        self._sheet_name : str | None = self._project.sheet_name
        self._run_id:str | None = None
        self._thread_id:str | None = None
        self._first_ai_reply = first_ai_reply
        self._queue =  queue.Queue()
        

    def upload_files(self):
        if self._file is None:
            print("No file to upload. Or Invaild column name")
            return

        if isinstance(self._file, BytesIO):
            file = AZURE_CLIENT.files.create(file=self._file, purpose="assistants")
        else:
            with open(self._file, "rb") as file_handle:
                file = AZURE_CLIENT.files.create(file=file_handle, purpose="assistants")

        self._file_id = file.id
        print(f"File uploaded successfully. File ID: {self._file_id}")


    def upload_file_to_assistant(self, response, filename):
        # Create a BytesIO object from the response content
        file_name , _ = os.path.splitext(filename)
        file_obj = BytesIO(response.content)
        file_obj.name = filename
     
        try:
            file_obj.seek(0)
            if file_obj.name.endswith(".csv"):
                df = pd.read_csv(file_obj)
                self._file = file_obj
            elif file_obj.name.endswith(".xlsx"):
                from openpyxl import load_workbook
                wb = load_workbook(file_obj, read_only=True) 
                actual_names = wb.sheetnames
                lookup = {name.strip() : name for name in actual_names}

                if self._sheet_name:
                    df = pd.read_excel(file_obj,engine="openpyxl",sheet_name=lookup[self._sheet_name])
                else:
                    df = pd.read_excel(file_obj,engine="openpyxl")
                buffer = BytesIO()
                df.to_csv(buffer)
                buffer.name = file_name + ".csv"
                self._file = buffer

            all_columns = set([c.strip() for c in df.columns])
            if self._user_column in all_columns:
                self._columns = all_columns
                return True

            return {"error": "Invaild Column name"}

        except ValueError as ve:
            print(f"ValueError: {ve}")
            raise


    def download_and_extract_files(self, file_ids):
        queryset = Files.objects.filter(id__in=file_ids).only("file")
        if not queryset:
            raise Exception("As your file(s) have been deleted, this chatbot only has access to the analysis but will not have access to your files.")

        uploaded_files = [file.file.name for file in queryset]
        for i in uploaded_files:
            filename = str(i).split("/")[-1]

            if filename.endswith(".csv") or filename.endswith(".xlsx"):
                _ , response = dbx.files_download(path=FOLDER_PATH + str(i))
                upload_file = self.upload_file_to_assistant(response, filename)

                if isinstance(upload_file, dict):
                    raise Exception(upload_file['error'])
            else:
                raise Exception("Invaild file format ,File must xlsx or CSV format")


    def create_assistant(self):
        assistant = AZURE_CLIENT.beta.assistants.create(
            instructions=f"""
            You analyze a CSV file attached for code_interpreter.

            ### Loading the data
            - Always load the CSV into a pandas DataFrame.
            - Inspect column names and data types before analysis.

            ### Identifying one-hot encoded columns
            - Automatically identify one-hot encoded columns as those that:
            - Have numeric dtype
            - Contain only values 0 or 1
            - Have column names containing an underscore (_), meaning Category_Subcategory

            ### When the user asks to plot
            - Identify the relevant category or categories by matching words in the user query to column names.
            - Select the corresponding one-hot encoded columns automatically.
            - Plot counts using matplotlib.
            - Always call plt.show() so the plot is returned as an image_file.

            ### When the user asks for analysis
            - Use the detected one-hot columns to filter rows.
            - Extract up to 10 text responses for illustrative examples.
            - Write a detailed explanation with examples.

            ### Open-ended response column
            - Detect the column containing free-text responses by choosing the column with string dtype and the highest average text length.


            ## Additional Instruction:
            - Follow the user instruction carefully to response queries better.
            - First analyze the user query and provide response in more deterministic way.
            
            """,
            model=OPENAI_MODEL_ENV1,
            tools=[{"type": "code_interpreter"}],
            tool_resources={"code_interpreter": {"file_ids": [self._file_id]}},
        )

        self._assistant_id = assistant.id


    def create_thread(self):
        threads = AZURE_CLIENT.beta.threads.create()
        
        if self._first_ai_reply:
            AZURE_CLIENT.beta.threads.messages.create(
            thread_id=threads.id,
            role="assistant",
            content=self._first_ai_reply
            )
           
    
        AZURE_CLIENT.beta.threads.messages.create(
            thread_id=threads.id,
            role="user",
            content=self._user_message,
            attachments=[{
                            "file_id": self._file_id,
                            "tools": [{"type": "code_interpreter"}],
                        }
                    ],
        )
        
        self._thread_id = threads.id
        
    def _get_instance(self):
        return  ExcelProjects.objects.get(id=self._project_id)
    
    def make_database_entry(self):
        instance = ExcelChatBotModel.objects.create(
                user=self._user,
                thread_id=self._thread_id,
                assistant_id=self._assistant_id,
                project_id=self._project_id,
                )
        return instance.id
        
    def add_message_to_thread(self):
        AZURE_CLIENT.beta.threads.messages.create(
            thread_id=self._thread_id, role="user", content=self._user_message
        )

    def run_new_chat(self):
        self.download_and_extract_files(self._project.file_ids)
        self.upload_files()
        self.create_assistant()
        self.create_thread()
        if not self._user.is_anonymous:
            instance_id = self.make_database_entry()
            return instance_id
        return None
        
    def generate_chat_title(self,instance_id, assistant_response):
        if instance_id:
            try:
                prompt = f"""
                Based on the text below, provide a descriptive chat title  Only output the title name, nothing else.
                {assistant_response}
                """
                response =AZURE_CLIENT.chat.completions.create(
                    model=OPENAI_MODEL_ENV1,
                    temperature=0,
                    messages=[
                        {"role": "system", "content": "You are a helpful assistant."},
                        {"role": "user", "content": prompt},
                    ],
                )
                project_name = response.choices[0].message.content
                name = project_name.replace('"', "")
            except Exception as e:
                name = "User chat"
            ExcelChatBotModel.objects.filter(id=instance_id).update(chat_title=name)
        
        
    def get_image(self,image_file_id):
        file_response = AZURE_CLIENT.files.content(image_file_id)
        image_bytes = file_response.read()
        image_base64 = base64.b64encode(image_bytes).decode("utf-8")
        return image_base64

       
    def run_functions(self, thread_id, assistant_id):
        instance_id = None
        try:
            if self._new_chat:
                instance_id = self.run_new_chat()  
            else:
                self._assistant_id = assistant_id
                self._thread_id = thread_id
                self.add_message_to_thread()

            def background_stream(self,new_chat,instance_id):
                event_handler=EventHandler(queue=self._queue)
                with AZURE_CLIENT.beta.threads.runs.stream(
                    thread_id=self._thread_id,
                    assistant_id=self._assistant_id,
                    event_handler=event_handler
                ) as stream:
                    stream.until_done()
                self._queue.put(None)
                if new_chat:
                    self.generate_chat_title(instance_id,event_handler._result)
                
            t = Thread(target=background_stream,args=(self,self._new_chat,instance_id), daemon=True)
            t.start()
            
            while True:
                msg = self._queue.get()
                if msg is None:
                    break
                base64_encoded_string = None
                if isinstance(msg, tuple) and msg[0] == "image_file_id":
                    _, image_file_id = msg
                    base64_encoded_string = self.get_image(image_file_id=image_file_id)
                    msg = None
                message_string = json.dumps(ExcelChatBotPayload(
                                error =  False,
                                error_message=None,
                                thread_id = self._thread_id,
                                assistant_id =  self._assistant_id,
                                text_response=msg,
                                image_data=base64_encoded_string
                            )
                           )
                yield f"data: {message_string}\n\n"
            message_string = json.dumps(ExcelChatBotPayload(
                                error =  False,
                                error_message=None,
                                thread_id = self._thread_id,
                                assistant_id =  self._assistant_id,
                                text_response="Response Done",
                                image_data=None
                            )
                           )
            yield f"data: {message_string}\n\n"
        except Exception as e:
            message_string =   json.dumps(ExcelChatBotPayload(
                error =  True,
                error_message=str(e),
                thread_id = self._thread_id,
                assistant_id =  self._assistant_id,
                text_response=None,
                image_data=None
            ))
            yield f"data: {message_string}\n\n"




class GeneralChatbot:
    def __init__(self,new_chat : bool , usermessage :str,user : object,files:list):
        self.new_chat = new_chat
        self.usermessage = usermessage
        self.user  = user
        self.thread_id = None
        self.assistant_id =  None
        self.vector_store_id = None
        self.files=files 
        self.files_stream = []
        self.queue = queue.Queue()
        
    def generate_chat_title(self,instance_id, assistant_response):
        try:
            prompt = f"""
            Based on the text below, provide a descriptive chat title  Only output the title name, nothing else.
            {assistant_response}
            """
            response =AZURE_CLIENT.chat.completions.create(
                model=OPENAI_MODEL_ENV1,
                temperature=0,
                messages=[
                    {"role": "system", "content": "You are a helpful assistant."},
                    {"role": "user", "content": prompt},
                ],
            )
            project_name = response.choices[0].message.content
            name = project_name.replace('"', "")
        except Exception as e:
            name = "User chat"
        GeneralChatbotModel.objects.filter(id=instance_id).update(chat_title=name)
        
        
    def upload_file_to_assistant(self, response, filename):
        filename , extension = os.path.splitext(filename)
        if extension.lower() in ['.xls','.xlsx']:
            extension = ".csv"
        file_obj = BytesIO(response.content)
        file_obj.name = filename+extension
        self.files_stream.append(file_obj)


    def download_and_extract_files(self):
        file_ids = [file['file_id'] for file in self.files]
        uploaded_files = chat_uploaded_files_filter_id(file_ids=file_ids)
        
        for i in uploaded_files:  
            try:
                filename = str(i).split("/")[-1]
                response = cache.get(i)
                if not response:
                    b, response = dbx.files_download(path=FOLDER_PATH + str(i))
                    cache.set(i, response, timeout=1500)
                self.upload_file_to_assistant(response, filename)
            except dropbox.exceptions.ApiError:
                raise Exception(f"{filename} file may deleted from backend. Please select another file.")
            
                
    def create_vector_store(self):
        if self.files:
            vector_store =  AZURE_CLIENT_1.vector_stores.create(
                    name=f"General_chatbot_{str(self.user)}", expires_after={"anchor": "last_active_at", "days": 100}
                )
            self.vector_store_id = vector_store.id
       
            self.download_and_extract_files()
            if self.files_stream:
                file_batch = AZURE_CLIENT_1.vector_stores.file_batches.upload_and_poll(
                    vector_store_id=vector_store.id, files=self.files_stream
                    )
                print(file_batch.status)
                print(file_batch.file_counts)
        
        
        
    def create_assistant(self):
        parameters = {"model":OPENAI_MODEL_5 ,"instructions" : "Act as chat Assistant","name":"General Chatbot"}

        if self.vector_store_id:
            parameters['tools'] =[{"type": "file_search"}]
            parameters['tool_resources'] =  {"file_search":{"vector_store_ids":[self.vector_store_id]}}

        assitant = AZURE_CLIENT_1.beta.assistants.create(**parameters)
        self.assistant_id = assitant.id
        
    def create_thread(self):
        thread = AZURE_CLIENT_1.beta.threads.create()
        self.thread_id = thread.id
    
    def add_message_to_thread(self):
        AZURE_CLIENT_1.beta.threads.messages.create(
            thread_id=self.thread_id,
            role="user",
            content = self.usermessage
        )
        
    def add_database_entry(self):
        instance = GeneralChatbotModel.objects.create(
            user = self.user,
            thread_id = self.thread_id,
            assistant_id = self.assistant_id
        )
        return instance.id
        
    def process_thread_message(self,new_chat , instance_id):
        event_handler = GeneralChatbotEventHandler(self.queue)
        with AZURE_CLIENT_1.beta.threads.runs.stream(thread_id=self.thread_id,assistant_id=self.assistant_id,event_handler=event_handler) as stream:
            stream.until_done()
        if new_chat:
            self.generate_chat_title(instance_id=instance_id,assistant_response=event_handler.result)
        print("Response Done")
       
    def run_functions(self,thread_id:str, assistant_id : str):
        instance_id = None
        try:
            if self.new_chat:
                self.create_vector_store()
                self.create_assistant()
                self.create_thread()
                if self.user.is_authenticated:
                    instance_id = self.add_database_entry()
            else:
                self.thread_id=thread_id
                self.assistant_id=assistant_id
            self.add_message_to_thread()   
            
            t = Thread(target=self.process_thread_message,args=(self.new_chat,instance_id), daemon=True)
            t.start()
        
            while True:
                msg = self.queue.get(timeout=90)
                if msg is None:
                    break
                message_string = json.dumps(ExcelChatBotPayload(
                                    error =  False,
                                    error_message=None,
                                    thread_id = self.thread_id,
                                    assistant_id =  self.assistant_id,
                                    text_response=msg,
                                )
                            )
                yield f"data: {message_string}\n\n"
            message_string = json.dumps(ExcelChatBotPayload(
                                error =  False,
                                error_message=None,
                                thread_id = self.thread_id,
                                assistant_id =  self.assistant_id,
                                text_response="Response Done",
                            )
                        )
            yield f"data: {message_string}\n\n"
        except Exception as e:
            message_string =   json.dumps(ExcelChatBotPayload(
                error =  True,
                error_message=str(e),
                thread_id = self.thread_id,
                assistant_id =  self.assistant_id,
                text_response=None,
            ))
            yield f"data: {message_string}\n\n"
        


